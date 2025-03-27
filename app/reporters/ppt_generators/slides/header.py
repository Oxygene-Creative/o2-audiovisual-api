from datetime import datetime
from pptx.util import Inches
from pptx.dml.color import RGBColor
from app.reporters.ppt_generators.slides.images import sanitize_image 
from app.reporters.ppt_generators.slides.colors import hex_to_rgb

def slide_header(report_date, data, slide, slide_width, logo_placeholder_idx):
    # Add title
    title = slide.shapes.title
    
    # title.text = f"Report For: {datetime.now().strftime('%d %B %Y')}"
    title.text = f"Report For: {report_date}"
    
    # Define the rectangle position and size
    left = 0
    top = 0
    width = slide_width 
    height = Inches(0.2) 
    
    # Add a rectangle shape spanning full width at the top
    shape = slide.shapes.add_shape(
        autoshape_type_id=1,  
        left=left,
        top=top,
        width=width,
        height=height
    )
    
    fill = shape.fill
    fill.solid()
    primary_rgb_color = hex_to_rgb(data['account']['brand_colors']['primary'])
    fill.fore_color.rgb = RGBColor(*primary_rgb_color)
    
    # set logo
    logo_placeholder = slide.placeholders[logo_placeholder_idx]
    logo_img = data['account']['logo']
    logo_img_stream = sanitize_image(logo_img)
    logo_placeholder.insert_picture(logo_img_stream)

