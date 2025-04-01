from app.reporters.ppt_generators.slides.colors import hex_to_rgb, change_text_color
from app.reporters.ppt_generators.slides.header import slide_header
from datetime import datetime
from app.reporters.charts.volume_chart import create_volume_chart_ppt
from app.reporters.charts.bar_chart import create_bar_chart_ppt
from pptx.enum.shapes import MSO_SHAPE_TYPE

def bind_common_placeholders(prs, slide, data, report_date, title="VOLUME OF MENTIONS"):
    primary_color_rgb = hex_to_rgb(data['account']['brand_colors']['primary'])
    text_color_rgb = hex_to_rgb("#666666")
    slide_header(report_date, data, slide, prs.slide_width, 12)
    
    heading = slide.placeholders[20] 
    heading.text = title.upper()
    change_text_color(heading, primary_color_rgb)
    
    summary = slide.placeholders[21] 
    summary.text = data['volume_summary']
    change_text_color(summary, text_color_rgb)
    
    mentions_summary = slide.placeholders[22] 
    mentions_summary.text = data['volume_mentions_summary']
    change_text_color(mentions_summary, text_color_rgb)
   
        
def add_volume_mentions_slide_area_chart(prs, data, date_format, report_date, title):
    try:
        
        # Add volume of mentions slide
        line_chart_layout = [layout for layout in prs.slide_layouts 
                        if layout.name == "Line Charts Slide"][0]
        slide = prs.slides.add_slide(line_chart_layout)
        
        # Bind some common placeholders
        bind_common_placeholders(prs, slide, data, report_date, title)
        
        volume_mentions_chart = slide.placeholders[23]
        if volume_mentions_chart.is_placeholder:
            
            dates = [datetime.strptime(date, "%Y-%m-%d") for date in data['volume_mentions']['dates']]
            
            create_volume_chart_ppt(
                volume_mentions_chart, 
                dates, 
                data['volume_mentions']['values'],
                date_format,
                data['account']['brand_colors']['primary']
            )

    except Exception as e:
        print(f"Error creating volume mentions slide: {e}")

def add_volume_mentions_slide_bar_chart(prs, data, report_date):
    try:
        # Add volume of mentions slide
        line_chart_layout = [layout for layout in prs.slide_layouts 
                        if layout.name == "Line Charts Slide"][0]
        slide = prs.slides.add_slide(line_chart_layout)
        
        # Bind some common placeholders
        bind_common_placeholders(prs, slide, data, report_date)
        
        volume_mentions_chart = slide.placeholders[23]
        if volume_mentions_chart.is_placeholder:
            bar_chart_colors = [hex_to_rgb(color) for color in data['volume_mentions']['colors']]
            create_bar_chart_ppt(
                volume_mentions_chart, 
                data['volume_mentions']['labels'], 
                data['volume_mentions']['values'],
                bar_chart_colors
            )

    except Exception as e:
        print(f"Error creating volume mentions slide: {e}")